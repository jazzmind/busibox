"""
PowerPoint builder and validator.

``build_presentation(spec, images)`` turns a :class:`PresentationSpec` into
``.pptx`` bytes with python-pptx. Every slide is drawn on the blank layout
with absolute geometry — a 16:9 canvas, a title band with an accent rule,
a footer with the deck title and slide number — so the result looks the
same in PowerPoint, Keynote and LibreOffice and does not depend on a
template's placeholders. Charts are native (editable in PowerPoint), tables
are real tables, images come from Busibox media (``render_chart`` output).

Verification re-opens the file, checks every slide has its title and the
figures/tables/charts it was meant to carry, flags slides too dense to
read, then renders through LibreOffice for a page count (which must equal
the slide count) and a first-slide thumbnail.
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from services.document_engine.office import OfficeToolError, pdf_first_page_png, pdf_page_count, soffice_convert, tool_availability
from services.document_engine.specs import PresentationSpec, SlideSpec, SourceSpec, ValidationIssue, ValidationReport

logger = logging.getLogger(__name__)

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# 16:9 canvas
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN
TITLE_TOP, TITLE_H = Inches(0.45), Inches(0.9)
BODY_TOP = Inches(1.55)
BODY_H = Inches(5.2)
FOOTER_TOP = Inches(6.95)

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2F, 0x54, 0x96)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xDC, 0xE6, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT = "Calibri"
HEAD_FONT = "Calibri Light"

DENSE_CHARS = 800
DENSE_BULLETS = 9
SOURCES_PER_SLIDE = 7

_SUB_RE = re.compile(r"^(?:-|•|\*)\s+")


@dataclass
class DeckBuildOutput:
    data: bytes
    report: ValidationReport
    summary: str
    thumbnail: Optional[bytes] = None
    stats: Dict[str, int] = field(default_factory=dict)


def collect_deck_image_ids(spec: PresentationSpec) -> List[str]:
    """Every Busibox file id the deck needs; the caller fetches these."""
    ids: List[str] = []
    for s in spec.slides:
        if s.image and s.image.file_id not in ids:
            ids.append(s.image.file_id)
    return ids


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------


def _textbox(slide, left, top, width, height, name: str):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return box


def _run(paragraph, text: str, size: int, *, bold=False, color=GREY, font=BODY_FONT, italic=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return run


def _title_band(slide, title: str, deck_title: str, index: int, total: int):
    box = _textbox(slide, MARGIN, TITLE_TOP, CONTENT_W, TITLE_H, "Title")
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    p = tf.paragraphs[0]
    _run(p, title, 30 if len(title) <= 60 else 24, bold=True, color=NAVY, font=HEAD_FONT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, TITLE_TOP + TITLE_H + Inches(0.02), CONTENT_W, Inches(0.03))
    rule.name = "Rule"
    rule.fill.solid()
    rule.fill.fore_color.rgb = NAVY
    rule.line.fill.background()
    _footer(slide, deck_title, index, total)


def _footer(slide, deck_title: str, index: int, total: int):
    left = _textbox(slide, MARGIN, FOOTER_TOP, CONTENT_W - Inches(1.2), Inches(0.35), "FooterTitle")
    _run(left.text_frame.paragraphs[0], deck_title[:90], 10, color=GREY)
    right = _textbox(slide, SLIDE_W - MARGIN - Inches(1.2), FOOTER_TOP, Inches(1.2), Inches(0.35), "SlideNumber")
    p = right.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{index} / {total}", 10, color=GREY)


def _bullet_size(bullets: List[str]) -> int:
    n = len(bullets)
    chars = sum(len(b) for b in bullets)
    if n <= 5 and chars <= 350:
        return 20
    if n <= 7 and chars <= 550:
        return 18
    if n <= 9 and chars <= 800:
        return 16
    return 14


def _bullets(slide, bullets: List[str], left, top, width, height, name: str, heading: Optional[str] = None, size: Optional[int] = None):
    box = _textbox(slide, left, top, width, height, name)
    tf = box.text_frame
    size = size or _bullet_size(bullets)
    first = True
    if heading:
        p = tf.paragraphs[0]
        _run(p, heading, size + 2, bold=True, color=BLUE)
        p.space_after = Pt(6)
        first = False
    for raw in bullets:
        level = 1 if _SUB_RE.match(raw) else 0
        text = _SUB_RE.sub("", raw, count=1)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(6 if level == 0 else 3)
        _run(p, ("• " if level == 0 else "– ") + text, size - (2 if level else 0), color=RGBColor(0x26, 0x26, 0x26))
    return box


def _picture(slide, data: bytes, left, top, max_w, max_h, name: str):
    pic = slide.shapes.add_picture(io.BytesIO(data), left, top)
    pic.name = name
    scale = min(max_w / pic.width, max_h / pic.height, 1.0 if pic.width <= max_w and pic.height <= max_h else 10.0)
    if scale != 1.0:
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
    # Centre in the box.
    pic.left = int(left + (max_w - pic.width) / 2)
    pic.top = int(top + (max_h - pic.height) / 2)
    return pic


def _caption(slide, text: str, left, top, width, name: str = "Caption"):
    box = _textbox(slide, left, top, width, Inches(0.4), name)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, text, 12, color=GREY, italic=True)


def _is_number(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and bool(re.match(r"^[\s$€£(]*-?[\d,]+(\.\d+)?%?[)\s]*$", v))


def _table(slide, headers: List[str], rows: List[List[Any]], left, top, width, height, name: str = "Table"):
    n_rows, n_cols = len(rows) + 1, len(headers)
    row_h = min(Inches(0.45), int(height / n_rows))
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, row_h * n_rows)
    shape.name = name
    table = shape.table
    size = 14 if n_rows <= 8 and n_cols <= 5 else 12 if n_rows <= 12 else 10
    # Column widths weighted by content length.
    weights = []
    for c in range(n_cols):
        longest = max([len(str(headers[c]))] + [len(str(r[c])) if c < len(r) and r[c] is not None else 0 for r in rows])
        weights.append(min(40, max(4, longest)))
    total = float(sum(weights))
    for c in range(n_cols):
        table.columns[c].width = int(width * weights[c] / total)
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = ""
        _run(cell.text_frame.paragraphs[0], str(h), size, bold=True, color=NAVY)
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT
    for r, row in enumerate(rows, start=1):
        for c in range(n_cols):
            v = row[c] if c < len(row) else None
            cell = table.cell(r, c)
            cell.text = ""
            text = "" if v is None else (f"{v:,.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else f"{v:,}" if isinstance(v, int) and not isinstance(v, bool) else str(v))
            p = cell.text_frame.paragraphs[0]
            if _is_number(v):
                p.alignment = PP_ALIGN.RIGHT
            _run(p, text, size, color=RGBColor(0x26, 0x26, 0x26))
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(0xF5, 0xF7, 0xFA)
    return shape


_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def _chart(slide, chart_spec, left, top, width, height, name: str = "Chart"):
    data = CategoryChartData()
    data.categories = list(chart_spec.categories)
    for srs in chart_spec.series:
        data.add_series(srs.name, list(srs.values))
    frame = slide.shapes.add_chart(_CHART_TYPES[chart_spec.type], left, top, width, height, data)
    frame.name = name
    chart = frame.chart
    chart.font.size = Pt(12)
    chart.font.name = BODY_FONT
    multi = len(chart_spec.series) > 1
    chart.has_legend = multi or chart_spec.type == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    if chart_spec.type == "pie":
        plot.has_data_labels = True
        plot.data_labels.number_format = chart_spec.number_format or "0%"
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
    else:
        if chart_spec.number_format:
            chart.value_axis.tick_labels.number_format = chart_spec.number_format
            chart.value_axis.tick_labels.number_format_is_linked = False
        chart.value_axis.has_major_gridlines = True
        if not multi and chart_spec.type != "line":
            plot.has_data_labels = True
            if chart_spec.number_format:
                plot.data_labels.number_format = chart_spec.number_format
                plot.data_labels.number_format_is_linked = False
    return frame


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------


def _title_slide(prs, title: str, subtitle: Optional[str], author: Optional[str], date_text: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.35))
    band.name = "Band"
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    box = _textbox(slide, MARGIN, Inches(2.2), CONTENT_W, Inches(1.8), "Title")
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    _run(tf.paragraphs[0], title, 40 if len(title) <= 50 else 32, bold=True, color=NAVY, font=HEAD_FONT)
    sub = _textbox(slide, MARGIN, Inches(4.1), CONTENT_W, Inches(0.8), "Subtitle")
    _run(sub.text_frame.paragraphs[0], subtitle or "", 20, color=GREY, font=HEAD_FONT)
    meta = _textbox(slide, MARGIN, Inches(5.1), CONTENT_W, Inches(0.6), "Meta")
    _run(meta.text_frame.paragraphs[0], " · ".join(x for x in (author, date_text) if x), 14, color=GREY)
    return slide


def _section_slide(prs, spec: SlideSpec, deck_title: str, index: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.name = "Background"
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    box = _textbox(slide, MARGIN, Inches(2.6), CONTENT_W, Inches(1.4), "Title")
    box.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    _run(box.text_frame.paragraphs[0], spec.title or "", 36, bold=True, color=WHITE, font=HEAD_FONT)
    if spec.subtitle:
        sub = _textbox(slide, MARGIN, Inches(4.1), CONTENT_W, Inches(0.9), "Subtitle")
        _run(sub.text_frame.paragraphs[0], spec.subtitle, 18, color=LIGHT, font=HEAD_FONT)
    num = _textbox(slide, SLIDE_W - MARGIN - Inches(1.2), FOOTER_TOP, Inches(1.2), Inches(0.35), "SlideNumber")
    num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _run(num.text_frame.paragraphs[0], f"{index} / {total}", 10, color=LIGHT)
    return slide


def _content_slide(prs, spec: SlideSpec, images: Dict[str, bytes], deck_title: str, index: int, total: int, issues: List[ValidationIssue]) -> Dict[str, int]:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title_band(slide, spec.title or "", deck_title, index, total)
    where = f"slide {index}"
    stats = {"pictures": 0, "tables": 0, "charts": 0}
    img = images.get(spec.image.file_id) if spec.image else None
    if spec.image and img is None:
        issues.append(ValidationIssue(severity="warning", location=where, message=f"image {spec.image.file_id} could not be embedded"))

    if spec.layout == "bullets":
        if img:
            _bullets(slide, spec.bullets, MARGIN, BODY_TOP, Inches(6.2), BODY_H, "Body")
            _picture(slide, img, MARGIN + Inches(6.5), BODY_TOP, CONTENT_W - Inches(6.5), BODY_H - Inches(0.5), "Picture")
            stats["pictures"] += 1
            if spec.image.caption:
                _caption(slide, spec.image.caption, MARGIN + Inches(6.5), BODY_TOP + BODY_H - Inches(0.45), CONTENT_W - Inches(6.5))
        else:
            _bullets(slide, spec.bullets, MARGIN, BODY_TOP, CONTENT_W, BODY_H, "Body")
    elif spec.layout == "two_column":
        col_w = (CONTENT_W - Inches(0.4)) / 2
        size = min(_bullet_size(spec.bullets), _bullet_size(spec.right_bullets))
        _bullets(slide, spec.bullets, MARGIN, BODY_TOP, col_w, BODY_H, "Body", spec.left_heading, size)
        _bullets(slide, spec.right_bullets, MARGIN + col_w + Inches(0.4), BODY_TOP, col_w, BODY_H, "BodyRight", spec.right_heading, size)
    elif spec.layout == "image":
        if img:
            cap_h = Inches(0.45) if spec.image.caption else 0
            _picture(slide, img, MARGIN, BODY_TOP, CONTENT_W, BODY_H - cap_h, "Picture")
            stats["pictures"] += 1
            if spec.image.caption:
                _caption(slide, spec.image.caption, MARGIN, BODY_TOP + BODY_H - Inches(0.45), CONTENT_W)
        elif spec.bullets:
            _bullets(slide, spec.bullets, MARGIN, BODY_TOP, CONTENT_W, BODY_H, "Body")
    elif spec.layout == "table":
        _table(slide, spec.table.headers, spec.table.rows, MARGIN, BODY_TOP, CONTENT_W, BODY_H, "Table")
        stats["tables"] += 1
    elif spec.layout == "chart":
        if spec.bullets:
            _chart(slide, spec.chart, MARGIN, BODY_TOP, Inches(7.6), BODY_H - Inches(0.4), "Chart")
            _bullets(slide, spec.bullets, MARGIN + Inches(7.9), BODY_TOP, CONTENT_W - Inches(7.9), BODY_H, "Body")
        else:
            _chart(slide, spec.chart, MARGIN, BODY_TOP, CONTENT_W, BODY_H - Inches(0.4), "Chart")
        stats["charts"] += 1
        if spec.chart.source:
            _caption(slide, f"Source: {spec.chart.source}", MARGIN, BODY_TOP + BODY_H - Inches(0.4), CONTENT_W, "Source")

    text_chars = sum(len(b) for b in spec.bullets) + sum(len(b) for b in spec.right_bullets)
    n_bullets = len(spec.bullets) + len(spec.right_bullets)
    if text_chars > DENSE_CHARS or n_bullets > DENSE_BULLETS:
        issues.append(ValidationIssue(severity="warning", location=where, message=f"dense slide ({n_bullets} bullets, {text_chars} characters) — consider splitting it or moving detail to the notes"))
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes
    return stats


def _sources_slides(prs, sources: List[SourceSpec], deck_title: str, start_index: int, total: int) -> int:
    chunks = [sources[i:i + SOURCES_PER_SLIDE] for i in range(0, len(sources), SOURCES_PER_SLIDE)]
    for k, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = "Sources" if len(chunks) == 1 else f"Sources ({k + 1}/{len(chunks)})"
        _title_band(slide, title, deck_title, start_index + k, total)
        box = _textbox(slide, MARGIN, BODY_TOP, CONTENT_W, BODY_H, "Body")
        tf = box.text_frame
        for i, src in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            n = k * SOURCES_PER_SLIDE + i + 1
            _run(p, f"{n}. {src.title}", 13, color=RGBColor(0x26, 0x26, 0x26))
            if src.url:
                _run(p, f"  {src.url}", 11, color=BLUE)
    return len(chunks)


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def build_pptx(spec: PresentationSpec, images: Dict[str, bytes], issues: List[ValidationIssue]) -> Tuple[bytes, Dict[str, int]]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    prs.core_properties.title = spec.title
    prs.core_properties.author = spec.author or "Busibox"
    prs.core_properties.subject = spec.subtitle or ""
    prs.core_properties.comments = "Generated by Busibox AI Chat"

    date_text = spec.date or dt.date.today().strftime("%B %d, %Y")
    body = list(spec.slides)
    add_title = spec.title_slide and (not body or body[0].layout != "title")
    n_sources = (len(spec.sources) + SOURCES_PER_SLIDE - 1) // SOURCES_PER_SLIDE if spec.sources else 0
    total = len(body) + (1 if add_title else 0) + n_sources

    index = 1
    if add_title:
        _title_slide(prs, spec.title, spec.subtitle, spec.author, date_text)
        index += 1
    stats = {"pictures": 0, "tables": 0, "charts": 0, "sections": 0}
    for s in body:
        if s.layout == "title":
            _title_slide(prs, s.title or spec.title, s.subtitle or spec.subtitle, spec.author, date_text)
        elif s.layout == "section":
            _section_slide(prs, s, spec.title, index, total)
            stats["sections"] += 1
        else:
            for k, v in _content_slide(prs, s, images, spec.title, index, total, issues).items():
                stats[k] += v
        if s.notes and s.layout in ("title", "section"):
            prs.slides[-1].notes_slide.notes_text_frame.text = s.notes
        index += 1
    if spec.sources:
        _sources_slides(prs, spec.sources, spec.title, index, total)

    buf = io.BytesIO()
    prs.save(buf)
    stats["slides"] = total
    return buf.getvalue(), stats


# ---------------------------------------------------------------------------
# Verify
# ---------------------------------------------------------------------------


def _verify_structure(data: bytes, spec: PresentationSpec, expected: Dict[str, int], images: Dict[str, bytes], issues: List[ValidationIssue], checks: List[str]) -> Dict[str, int]:
    prs = Presentation(io.BytesIO(data))
    n = len(prs.slides)
    if n != expected["slides"]:
        issues.append(ValidationIssue(severity="error", location="deck", message=f"expected {expected['slides']} slides, file has {n}"))
    pictures = tables = charts = untitled = 0
    for i, slide in enumerate(prs.slides, start=1):
        title = next((sh for sh in slide.shapes if sh.name == "Title"), None)
        if title is None or not title.text_frame.text.strip():
            untitled += 1
            issues.append(ValidationIssue(severity="error", location=f"slide {i}", message="slide has no title"))
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                tables += 1
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
    wanted_pictures = sum(1 for s in spec.slides if s.image and s.image.file_id in images)
    if pictures < wanted_pictures:
        issues.append(ValidationIssue(severity="error", location="deck", message=f"{wanted_pictures - pictures} image(s) did not embed"))
    if tables != sum(1 for s in spec.slides if s.layout == "table"):
        issues.append(ValidationIssue(severity="error", location="deck", message="a table slide has no table"))
    if charts != sum(1 for s in spec.slides if s.layout == "chart"):
        issues.append(ValidationIssue(severity="error", location="deck", message="a chart slide has no chart"))
    checks.append(f"re-opened the deck: {n} slide(s), {pictures} picture(s), {tables} table(s), {charts} chart(s)" + (f", {untitled} untitled" if untitled else ""))
    return {"slides": n, "pictures": pictures, "tables": tables, "charts": charts}


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def build_presentation(spec: PresentationSpec, images: Optional[Dict[str, bytes]] = None, timeout: int = 90, thumbnail: bool = True) -> DeckBuildOutput:
    """Build and verify a deck. LibreOffice problems degrade to warnings;
    a slide count mismatch or a missing title/figure is an error."""
    images = images or {}
    issues: List[ValidationIssue] = []
    checks: List[str] = []
    data, expected = build_pptx(spec, images, issues)
    checks.append(f"built {expected['slides']} slide(s) with python-pptx")
    stats = _verify_structure(data, spec, expected, images, issues, checks)

    pages: Optional[int] = None
    thumb: Optional[bytes] = None
    tools = tool_availability()
    if tools.soffice:
        with tempfile.TemporaryDirectory(prefix="docgen-pptx-") as tmp:
            path = Path(tmp) / "deck.pptx"
            path.write_bytes(data)
            try:
                pdf = soffice_convert(path, "pdf", timeout=timeout)
                pages = pdf_page_count(pdf)
                checks.append(f"rendered to PDF with LibreOffice — {pages if pages is not None else '?'} page(s)")
                if pages is not None and pages != stats["slides"]:
                    issues.append(ValidationIssue(severity="error", location="deck", message=f"LibreOffice rendered {pages} page(s) for {stats['slides']} slide(s)"))
                if thumbnail:
                    thumb = pdf_first_page_png(pdf)
                    if thumb is None:
                        issues.append(ValidationIssue(severity="warning", location="deck", message="first-slide thumbnail could not be rendered"))
            except OfficeToolError as exc:
                issues.append(ValidationIssue(severity="warning", location="deck", message=f"PDF render skipped: {exc} (is libreoffice-impress installed?)"))
    else:
        issues.append(ValidationIssue(severity="warning", location="deck", message="LibreOffice is not installed; render check and thumbnail unavailable"))

    ok = not any(i.severity == "error" for i in issues)
    report = ValidationReport(ok=ok, checks=checks, issues=issues, pages=pages)
    bits = [f"{stats['slides']} slide(s)"]
    if stats["charts"]:
        bits.append(f"{stats['charts']} chart(s)")
    if stats["tables"]:
        bits.append(f"{stats['tables']} table(s)")
    if stats["pictures"]:
        bits.append(f"{stats['pictures']} image(s)")
    if spec.sources:
        bits.append(f"{len(spec.sources)} source(s)")
    dense = sum(1 for i in issues if "dense slide" in i.message)
    summary = f"{spec.filename}: " + ", ".join(bits) + ("." if ok else f". {sum(1 for i in issues if i.severity == 'error')} error(s).") + (f" {dense} dense slide(s) flagged." if dense else "")
    return DeckBuildOutput(data=data, report=report, summary=summary, thumbnail=thumb, stats=stats)
