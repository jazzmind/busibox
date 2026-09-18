"""Tables in Word and PowerPoint files must end up in the extracted text.

Regression test for tabular content being dropped at ingest: ``_extract_docx``
built its text from ``doc.paragraphs`` (which excludes paragraphs inside table
cells) and ``_extract_pptx`` skipped table shapes (no ``.text``). Table content
only ever reached ``ExtractionResult.tables``, which the worker counts but never
chunks or embeds, so a rate table or an appendix table was unsearchable.

Run with::

    make test-docker SERVICE=data ARGS="tests/unit/test_office_table_extraction.py"
"""

from __future__ import annotations

from pathlib import Path

import pytest

from processors.text_extractor import TextExtractor, table_rows_to_markdown

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def extractor(tmp_path: Path) -> TextExtractor:
    return TextExtractor({"temp_dir": str(tmp_path / "work"), "marker_enabled": False})


# ---------------------------------------------------------------------------
# Markdown rendering
# ---------------------------------------------------------------------------


def test_table_rows_to_markdown_pads_escapes_and_drops_empty_rows():
    rows = [
        ["State", "Program", "Max Weekly Benefit (2026)"],
        ["", "", ""],
        ["Massachusetts", "MA  PFML", "$1,230.39"],
        ["New York", "DBL / PFL"],  # ragged row
        ["Rhode Island", "TDI | TCI", "$1,103"],
    ]
    md = table_rows_to_markdown(rows)
    lines = md.splitlines()
    assert lines[0] == "| State | Program | Max Weekly Benefit (2026) |"
    assert lines[1] == "|---|---|---|"
    assert "| Massachusetts | MA PFML | $1,230.39 |" in lines  # whitespace collapsed
    assert "| New York | DBL / PFL |  |" in lines  # padded to table width
    assert "| Rhode Island | TDI \\| TCI | $1,103 |" in lines  # pipe escaped
    assert len(lines) == 5  # header, rule, three data rows (empty row dropped)


def test_table_rows_to_markdown_empty_table_is_empty_string():
    assert table_rows_to_markdown([]) == ""
    assert table_rows_to_markdown([["", None], [" "]]) == ""


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _build_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Massachusetts Paid Family and Medical Leave", level=1)
    doc.add_paragraph("Reasons you can take leave, and for how long:")
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Reason", "Maximum Duration"
    t.cell(1, 0).text, t.cell(1, 1).text = "Bonding with a new child", "Up to 12 weeks"
    t.cell(2, 0).text, t.cell(2, 1).text = "Your own serious health condition", "Up to 20 weeks"
    doc.add_paragraph("Leave can be taken continuously or intermittently.")
    doc.add_heading("Appendix: State Paid Leave Programs at a Glance", level=1)
    t2 = doc.add_table(rows=3, cols=3)
    t2.cell(0, 0).text, t2.cell(0, 1).text, t2.cell(0, 2).text = "State", "Program", "Max Weekly Benefit"
    t2.cell(1, 0).text, t2.cell(1, 1).text, t2.cell(1, 2).text = "Colorado", "FAMLI", "$1,381.45"
    # Horizontally merged cell in the last row: must appear once, not three times.
    merged = t2.cell(2, 0).merge(t2.cell(2, 2))
    merged.text = "Maryland's program starts in 2027-28."
    doc.add_paragraph("Sources: state agency notices.")
    doc.save(str(path))


def test_docx_tables_are_in_text_in_document_order(extractor: TextExtractor, tmp_path: Path):
    path = tmp_path / "guide.docx"
    _build_docx(path)

    result = extractor.extract(str(path), DOCX_MIME)
    text = result.text

    # Table content is now part of the embedded text …
    assert "| Bonding with a new child | Up to 12 weeks |" in text
    assert "| Colorado | FAMLI | $1,381.45 |" in text
    # … placed where the table sits in the document, not appended at the end.
    assert text.index("for how long:") < text.index("| Reason | Maximum Duration |") < text.index("continuously")
    assert text.index("Appendix: State Paid Leave") < text.index("| State | Program |") < text.index("Sources:")

    # Merged cells are emitted once.
    assert text.count("Maryland's program starts") == 1
    assert "| Maryland's program starts in 2027-28. |" in text

    # Structured tables are still returned for metadata, with merged cells collapsed.
    assert len(result.tables) == 2
    assert result.tables[0]["data"][1] == ["Bonding with a new child", "Up to 12 weeks"]
    assert result.tables[1]["data"][2] == ["Maryland's program starts in 2027-28."]
    assert result.metadata["extraction_method"] == "python-docx"
    assert result.metadata["table_count"] == 2


def test_docx_without_tables_is_unchanged(extractor: TextExtractor, tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("")
    doc.add_paragraph("Second paragraph.")
    path = tmp_path / "plain.docx"
    doc.save(str(path))

    result = extractor.extract(str(path), DOCX_MIME)
    assert result.text == "First paragraph.\nSecond paragraph."
    assert result.tables == []


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "2026 Contribution Rates"
    shape = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    table = shape.table
    table.cell(0, 0).text, table.cell(0, 1).text = "Share", "Rate"
    table.cell(1, 0).text, table.cell(1, 1).text = "Employee", "0.46%"
    table.cell(2, 0).text, table.cell(2, 1).text = "Employer", "0.42%"
    prs.save(str(path))


def test_pptx_table_shapes_are_in_text(extractor: TextExtractor, tmp_path: Path):
    path = tmp_path / "rates.pptx"
    _build_pptx(path)

    result = extractor.extract(str(path), PPTX_MIME)
    assert "=== Slide 1 ===" in result.text
    assert "2026 Contribution Rates" in result.text
    assert "| Share | Rate |" in result.text
    assert "| Employee | 0.46% |" in result.text
    assert "| Employer | 0.42% |" in result.text
    assert result.page_count == 1
